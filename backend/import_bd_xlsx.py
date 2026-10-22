#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Importador de Base de Datos Institucional MIPPCI (bd.xlsx)
Combina y anexa los 1,122 funcionarios reales con los datos de ejemplo preexistentes.
"""

import os
import sys
import re
import json
import uuid
import sqlite3
import unicodedata
from datetime import datetime
import openpyxl

# Rutas del proyecto
if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
EXCEL_PATH = os.path.join(BASE_DIR, 'bd.xlsx')
SQLITE_PATH = os.path.join(BASE_DIR, 'backend', 'combos_platform.sqlite')
SEED_JSON_PATH = os.path.join(BASE_DIR, 'backend', 'production_seed.json')

def clean_mojibake(text):
    if not text or not isinstance(text, str):
        return ''
    # Corregir secuencias comunes de codificación corrupta
    replacements = {
        'Ã“': 'Ó', 'Ã‰': 'É', 'Ã\x81': 'Á', 'Ã\x8d': 'Í', 'Ãš': 'Ú', 'Ã‘': 'Ñ',
        'Ã³': 'ó', 'Ã©': 'é', 'Ã¡': 'á', 'Ã\xad': 'í', 'Ãº': 'ú', 'Ã±': 'ñ',
        'NÂ°': 'N°', 'N ': 'N° ', 'N°': 'N°', '  ': ' '
    }
    for k, v in replacements.items():
        text = text.replace(k, v)
    return text.strip()

def clean_name(raw):
    if not raw:
        return ''
    text = clean_mojibake(str(raw)).strip()
    if ',' in text:
        parts = text.split(',')
        text = f"{parts[1].strip()} {parts[0].strip()}"
    words = text.split()
    lower_words = {'de', 'del', 'la', 'las', 'los', 'y', 'e', 'da', 'di'}
    capitalized = []
    for i, w in enumerate(words):
        w_lower = w.lower()
        if i > 0 and w_lower in lower_words:
            capitalized.append(w_lower)
        else:
            capitalized.append(w.capitalize())
    return ' '.join(capitalized)

def strip_accents(text):
    return ''.join(c for c in unicodedata.normalize('NFD', text) if unicodedata.category(c) != 'Mn')

def generate_email(full_name, doc_digits, used_emails):
    clean = strip_accents(full_name).lower()
    words = re.findall(r'[a-z0-9]+', clean)
    if len(words) >= 2:
        base = f"{words[0]}.{words[-1]}"
    elif len(words) == 1:
        base = words[0]
    else:
        base = f"user{doc_digits}"
    
    email = f"{base}@mippci.gob.ve"
    if email in used_emails:
        email = f"{base}.{doc_digits}@mippci.gob.ve"
    
    counter = 1
    while email in used_emails:
        email = f"{base}{counter}.{doc_digits}@mippci.gob.ve"
        counter += 1
        
    used_emails.add(email)
    return email

def main():
    print("=" * 70)
    print("🚀 INICIANDO INGESTA DE bd.xlsx A LA BASE DE DATOS INSTITUCIONAL")
    print("=" * 70)

    if not os.path.exists(EXCEL_PATH):
        print(f"❌ Error: Archivo no encontrado en {EXCEL_PATH}")
        sys.exit(1)

    # 1. Conectar a SQLite
    conn = sqlite3.connect(SQLITE_PATH)
    cursor = conn.cursor()

    # Migración de columnas opcionales en SQLite si faltan
    try:
        cursor.execute("ALTER TABLE workers ADD COLUMN job_title TEXT DEFAULT ''")
    except sqlite3.OperationalError:
        pass
    try:
        cursor.execute("ALTER TABLE workers ADD COLUMN hire_date TEXT DEFAULT ''")
    except sqlite3.OperationalError:
        pass
    conn.commit()

    # 2. Cargar departamentos existentes
    cursor.execute("SELECT id, name FROM departments")
    existing_depts = {row[1].strip().lower(): row[0] for row in cursor.fetchall()}

    # Cargar emails existentes para evitar colisiones
    cursor.execute("SELECT email FROM workers")
    used_emails = {row[0].strip().lower() for row in cursor.fetchall() if row[0]}

    # Cargar trabajadores existentes indexados por número de cédula (dígitos)
    cursor.execute("SELECT id, document_id, full_name, role FROM workers")
    existing_workers_by_digits = {}
    for wid, doc_id, name, role in cursor.fetchall():
        digits = re.sub(r'[^0-9]', '', doc_id)
        existing_workers_by_digits[digits] = {
            'id': wid,
            'document_id': doc_id,
            'full_name': name,
            'role': role
        }

    print(f"📊 Registros previos en BD: {len(existing_workers_by_digits)} funcionarios")
    print(f"🏢 Departamentos previos en BD: {len(existing_depts)}")

    # 3. Leer Excel bd.xlsx
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    sheet_name = 'bd' if 'bd' in wb.sheetnames else wb.sheetnames[0]
    sheet = wb[sheet_name]

    rows = list(sheet.iter_rows(values_only=True))
    header = rows[0]
    data_rows = rows[1:]
    print(f"📄 Archivo bd.xlsx leído: {len(data_rows)} filas de funcionarios.")

    inserted_count = 0
    updated_count = 0
    dept_created_count = 0

    now_iso = datetime.now().isoformat()
    today_str = datetime.now().strftime('%Y-%m-%d')

    for r in data_rows:
        raw_ced = r[0]
        if raw_ced is None:
            continue
        ced_digits = re.sub(r'[^0-9]', '', str(raw_ced))
        if not ced_digits:
            continue

        raw_name = r[1]
        raw_date = r[2]
        raw_unit = r[3]
        raw_cargo = r[4]

        # Limpieza de datos
        name_clean = clean_name(raw_name)
        unit_clean = clean_mojibake(str(raw_unit or 'MIPPCI Sede Principal')).strip()
        cargo_clean = clean_mojibake(str(raw_cargo or 'Funcionario')).strip()

        # Formato de fecha
        hire_date_str = ''
        if isinstance(raw_date, datetime):
            hire_date_str = raw_date.strftime('%Y-%m-%d')
        elif raw_date:
            hire_date_str = str(raw_date).strip()[:10]

        # Resolver departamento
        dept_key = unit_clean.lower()
        dept_id = existing_depts.get(dept_key)
        if not dept_id:
            dept_id = str(uuid.uuid4())
            dept_code = 'DEP_' + str(uuid.uuid4())[:6].upper()
            cursor.execute(
                "INSERT INTO departments (id, code, name, cost_center, created_at) VALUES (?, ?, ?, ?, ?)",
                (dept_id, dept_code, unit_clean, 'CC-MIPPCI', now_iso)
            )
            existing_depts[dept_key] = dept_id
            dept_created_count += 1

        # Verificar si ya existe en BD
        existing = existing_workers_by_digits.get(ced_digits)
        if existing:
            # Actualizar datos complementarios pero preservar ID, rol (ej: admin/supervisor) y credenciales
            wid = existing['id']
            # Mantener document_id tal como estaba si ya existía (ej: 13537582)
            cursor.execute("""
                UPDATE workers 
                SET department_id = ?, job_title = ?, hire_date = ?, updated_at = ?
                WHERE id = ?
            """, (dept_id, cargo_clean, hire_date_str, now_iso, wid))
            updated_count += 1
        else:
            # Insertar nuevo funcionario
            wid = str(uuid.uuid4())
            doc_id = f"V-{ced_digits}"
            email = generate_email(name_clean, ced_digits, used_emails)

            cursor.execute("""
                INSERT INTO workers (
                    id, document_id, full_name, email, department_id,
                    shift, status, worker_type, role, has_smartphone,
                    contingency_reason, phone_number, avatar_url,
                    created_at, updated_at, job_title, hire_date
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                wid, doc_id, name_clean, email, dept_id,
                'jornada_continua', 'active', 'activo', 'worker', 1,
                None, None, None,
                now_iso, now_iso, cargo_clean, hire_date_str
            ))
            existing_workers_by_digits[ced_digits] = {
                'id': wid,
                'document_id': doc_id,
                'full_name': name_clean,
                'role': 'worker'
            }
            inserted_count += 1

    conn.commit()

    print(f"✅ Inserción completa en SQLite:")
    print(f"   - Funcionarios nuevos creados: {inserted_count}")
    print(f"   - Funcionarios existentes actualizados: {updated_count}")
    print(f"   - Nuevas unidades / departamentos creados: {dept_created_count}")

    # 4. Asignar Daily Entitlements automáticos para la jornada de hoy (Estación 1 y Estación 2)
    print("\n📦 Garantizando asignaciones de combos (Víveres y Proteína) para hoy...")
    cursor.execute("SELECT id FROM combos WHERE category IN ('bolsa_viveres', 'almuerzo') OR name LIKE '%Víveres%' ORDER BY id ASC LIMIT 1")
    v_row = cursor.fetchone()
    v_combo_id = v_row[0] if v_row else 'c1a2b3c4-1111-2222-3333-444455556666'

    cursor.execute("SELECT id FROM combos WHERE category IN ('combo_proteico', 'proteico') OR name LIKE '%Prote%' ORDER BY id ASC LIMIT 1")
    p_row = cursor.fetchone()
    p_combo_id = p_row[0] if p_row else 'c1a2b3c4-1111-2222-3333-444455557777'

    cursor.execute("SELECT id FROM workers WHERE status = 'active'")
    all_workers = cursor.fetchall()

    ent_count = 0
    for (w_id,) in all_workers:
        # Estacion 1
        try:
            cursor.execute("""
                INSERT OR IGNORE INTO daily_entitlements (id, worker_id, combo_id, valid_date, shift, status)
                VALUES (?, ?, ?, ?, 'estacion_1', 'pending')
            """, (str(uuid.uuid4()), w_id, v_combo_id, today_str))
            if cursor.rowcount > 0:
                ent_count += 1
        except Exception:
            pass

        # Estacion 2
        try:
            cursor.execute("""
                INSERT OR IGNORE INTO daily_entitlements (id, worker_id, combo_id, valid_date, shift, status)
                VALUES (?, ?, ?, ?, 'estacion_2', 'pending')
            """, (str(uuid.uuid4()), w_id, p_combo_id, today_str))
            if cursor.rowcount > 0:
                ent_count += 1
        except Exception:
            pass

    conn.commit()
    print(f"✅ Asignaciones diarias verificadas/agregadas: {ent_count} combos listos para retiro hoy.")

    # 5. Actualizar backend/production_seed.json
    print("\n💾 Actualizando backend/production_seed.json...")
    cursor.execute("SELECT id, code, name, cost_center, created_at FROM departments")
    all_depts_db = [
        {'id': r[0], 'code': r[1], 'name': r[2], 'cost_center': r[3], 'created_at': r[4]}
        for r in cursor.fetchall()
    ]

    cursor.execute("""
        SELECT id, document_id, full_name, email, department_id, shift, status,
               worker_type, role, has_smartphone, contingency_reason, phone_number,
               avatar_url, created_at, updated_at
        FROM workers
    """)
    all_workers_db = [
        {
            'id': r[0],
            'document_id': r[1],
            'full_name': r[2],
            'email': r[3],
            'department_id': r[4],
            'shift': r[5],
            'status': r[6],
            'worker_type': r[7],
            'role': r[8],
            'has_smartphone': r[9],
            'contingency_reason': r[10],
            'phone_number': r[11],
            'avatar_url': r[12],
            'created_at': r[13],
            'updated_at': r[14]
        }
        for r in cursor.fetchall()
    ]

    seed_data = {}
    if os.path.exists(SEED_JSON_PATH):
        try:
            with open(SEED_JSON_PATH, 'r', encoding='utf-8') as f:
                seed_data = json.load(f)
        except Exception as e:
            print("Aviso leyendo seed previo:", e)

    seed_data['departments'] = all_depts_db
    seed_data['workers'] = all_workers_db

    with open(SEED_JSON_PATH, 'w', encoding='utf-8') as f:
        json.dump(seed_data, f, ensure_ascii=False, indent=2)

    print(f"✅ production_seed.json sincronizado con {len(all_workers_db)} funcionarios y {len(all_depts_db)} departamentos.")

    # Total final
    cursor.execute("SELECT count(*) FROM workers")
    total_workers_db = cursor.fetchone()[0]
    cursor.execute("SELECT count(*) FROM daily_entitlements")
    total_entitlements_db = cursor.fetchone()[0]

    conn.close()

    print("\n" + "=" * 70)
    print(f"🎉 INGESTA EXITOSA:")
    print(f"   - Total de Trabajadores en Base de Datos: {total_workers_db}")
    print(f"   - Total de Asignaciones Diarias Activas:  {total_entitlements_db}")
    print("=" * 70)

if __name__ == '__main__':
    main()
